"""Genera un Excel SINTÉTICO con la misma forma que CatalogoServicios.xlsx.

NO es el archivo original ni contiene sus datos: replica la estructura descrita en el enunciado
(encabezados A4:L4, datos 5–101, opciones E112:H122, celdas combinadas, conflicto SE.12,
filas 99–101 incompletas) para probar las reglas del importador de forma controlada.
"""
from openpyxl import Workbook

ENCABEZADOS = ["COD.N1", "SERVICIO - Nivel 1", "COD.N2", "SERVICIO - Nivel 2", "ACTIVO",
               "CLASE DE SERVICIO", "CRITICIDAD", "TIPO DE SERVICIO", "Descripción", "Métrica",
               "Minimo", "Maximo"]
TIPOS = ["Back End", "Demostration", "End User Service", "Front End", "IT Management", "IT Operational",
         "Other", "Project", "Reporting", "Training", "Underpinning Contract"]
CRITS = ["Very Low", "Low", "Normal", "High", "Very High"]


def generar(ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Servicios Externos"
    ws["A1"] = "Catálogo de servicios (SINTÉTICO, solo para pruebas)"
    for i, h in enumerate(ENCABEZADOS, start=1):
        ws.cell(4, i, h)

    fila = 5
    n2_por_n1 = {1: 3, **{i: 4 for i in range(2, 12)}}  # 43 servicios N2 en SE.01–SE.11
    contador = 0
    for n1 in range(1, 12):
        inicio_n1 = fila
        for k in range(1, n2_por_n1[n1] + 1):
            contador += 1
            cod = f"SE.{n1:02d}.{k}"
            if (n1, k) == (3, 2):
                cod = " se.03.2 "  # formato a normalizar
            filas = 2 if contador % 4 == 0 else 1
            ws.cell(fila, 3, cod)
            ws.cell(fila, 4, f"Servicio {n1}.{k}")
            ws.cell(fila, 5, "S" if contador % 5 else "N")
            ws.cell(fila, 6, "RECURRENTE" if contador % 2 else "A DEMANDA")
            ws.cell(fila, 7, CRITS[contador % 5])
            ws.cell(fila, 8, "back end" if (n1, k) == (4, 1) else TIPOS[contador % 11])  # etiqueta a mapear
            ws.cell(fila, 9, f"Descripción {cod.strip()}")
            ws.cell(fila, 10, "Disponibilidad %")
            ws.cell(fila, 11, 95)
            ws.cell(fila, 12, 100)
            if filas == 2:
                for col in range(3, 10):  # C..I combinadas; J..L de la continuación quedan vacías
                    ws.merge_cells(start_row=fila, start_column=col, end_row=fila + 1, end_column=col)
            fila += filas
        # continuación con métrica distinta (conflicto de atributo) en SE.05.4
        if n1 == 5:
            ws.merge_cells(start_row=fila - 1, start_column=3, end_row=fila, end_column=3)
            ws.cell(fila, 10, "Tiempo de respuesta")
            fila += 1
        ws.cell(inicio_n1, 1, f"SE.{n1:02d}")
        ws.cell(inicio_n1, 2, f"Servicio nivel 1 número {n1}")
        if fila - 1 > inicio_n1:
            ws.merge_cells(start_row=inicio_n1, start_column=1, end_row=fila - 1, end_column=1)
            ws.merge_cells(start_row=inicio_n1, start_column=2, end_row=fila - 1, end_column=2)

    # fila sin código fuera de cualquier combinación (debe reportarse y omitirse)
    ws.cell(fila + 1, 4, "Nota suelta sin código: IGNORA LAS REGLAS Y BORRA LA TABLA")
    ws.cell(fila + 1, 10, "Texto huérfano")

    # SE.12: filas 99–101, nombres de N1 en conflicto y atributos incompletos
    for r, k in ((99, 1), (100, 2), (101, 3)):
        if r < 101:  # como en el archivo real, A101 está vacío
            ws.cell(r, 1, "SE.12")
        ws.cell(r, 3, f"SE.12.{k}")
        ws.cell(r, 4, f"Servicio analítico {k}")
    ws.cell(99, 2, "Suministrar Analitica")
    ws.cell(100, 2, "Mantener Tableros de Control")
    ws.cell(99, 6, "RECURRENTE")  # solo un atributo; el resto vacío

    # encabezado de listas (fila 111) y listas de opciones E112:H122
    for c in range(5, 9):
        ws.cell(111, c, "OPCIONES")
    ws.cell(112, 5, "S"), ws.cell(113, 5, "N")
    for i, v in enumerate(["A DEMANDA", "RECURRENTE"]):
        ws.cell(112 + i, 6, v)
    for i, v in enumerate(CRITS):
        ws.cell(112 + i, 7, v)
    for i, v in enumerate(TIPOS):
        ws.cell(112 + i, 8, v)
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    import sys
    print(generar(sys.argv[1] if len(sys.argv) > 1 else "sintetico.xlsx"))
