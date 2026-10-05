"""Importador de la hoja "Servicios Externos" de CatalogoServicios.xlsx.

Reglas (documentadas en docs/RESOLUCION.md §4 y docs/contexto/analisis-excel.md):
R1  Celdas combinadas: el valor de cualquier celda dentro de un rango combinado es el de la
    celda principal (superior izquierda) de ESE rango. Nunca se propaga fuera del rango.
R2  Un servicio N2 = el rango combinado de la columna C (o la fila sola si C no está combinada).
    Las demás filas del rango son "filas de continuación" del mismo servicio.
R3  Una fila sin código N2 que no pertenece a ningún rango de C y tiene contenido propio se
    reporta (FILA_SIN_CODIGO) y se omite; no se asigna a otro servicio.
R4  Conflicto de nombre de un N1 (p. ej. SE.12): se usa NOMBRE_CANONICO_N1 si está definido,
    si no la primera aparición; todos los nombres se guardan en nombres_alternativos y se
    emite CONFLICTO_NOMBRE_N1.
R5  Conflicto de atributo dentro de un N2 (filas de continuación con valores distintos):
    se conserva el primer valor no vacío (fila superior) y se emite CONFLICTO_ATRIBUTO.
R6  Los códigos se tratan como texto. Si la normalización (espacios/mayúsculas) cambia el
    valor, se guarda codigo_original y se emite CODIGO_NORMALIZADO.
R7  Atributos vacíos o fuera de catálogo no se inventan: quedan NULL, se guarda el valor
    original y el servicio queda en estado_revision='REVISION'.
R9  N2 sin COD.N1 (caso real: SE.12.3, fila 101): se asocia al N1 indicado por el prefijo de su
    propio código (SE.12.3 -> SE.12) solo si ese N1 existe, y se emite N1_DERIVADO_DEL_CODIGO.
R8  Idempotencia: la llave es el código. Se guarda un hash de los valores de origen; si el
    hash no cambió el registro no se toca (respeta ediciones hechas en la aplicación).
"""
import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import config
from .models import (ClaseServicio, Criticidad, Importacion, MapeoEtiqueta, ObservacionImportacion,
                     ServicioN1, ServicioN2, TipoServicio)

FILA_ENCABEZADO = 4
FILA_INICIO, FILA_FIN = 5, 101
OPCIONES = (112, 122, 5, 8)  # E112:H122
COLUMNAS = ["COD.N1", "SERVICIO - Nivel 1", "COD.N2", "SERVICIO - Nivel 2", "ACTIVO",
            "CLASE DE SERVICIO", "CRITICIDAD", "TIPO DE SERVICIO", "Descripción", "Métrica",
            "Minimo", "Maximo"]
CAMPOS = ["cod_n1", "nombre_n1", "cod_n2", "nombre_n2", "activo", "clase", "criticidad", "tipo",
          "descripcion", "metrica", "minimo", "maximo"]

# R4 — decisión documentada para el conflicto de SE.12 (filas 99 y 100).
# "Suministrar Analitica" es el nombre más general: abarca a "Mantener Tableros de Control",
# que describe una actividad concreta (propia de un nivel 2). Además es la primera aparición.
NOMBRE_CANONICO_N1 = {"SE.12": "Suministrar Analitica"}

# Equivalencias explícitas de ACTIVO (todo lo demás se conserva como desconocido)
ACTIVO_ALIAS = {"S": "S", "N": "N", "SI": "S", "NO": "N"}


def _sin_tildes(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def clave(s: str) -> str:
    """Forma comparable de una etiqueta: sin tildes, minúsculas, espacios colapsados."""
    return re.sub(r"\s+", " ", _sin_tildes(str(s))).strip().casefold()


def texto(v) -> str | None:
    """Convierte el valor de celda a texto sin perder formato de códigos; vacío -> None."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def normalizar_codigo(raw: str) -> str:
    return re.sub(r"\s+", "", raw).upper()


@dataclass
class Celda:
    valor: str | None
    origen: str  # coordenada de la celda que aporta el valor (principal del rango si es combinada)
    combinada: str | None  # rango combinado, si aplica


@dataclass
class Registro:
    codigo: str
    codigo_original: str
    filas: list[int] = field(default_factory=list)
    valores: dict[str, list[tuple[int, str, str]]] = field(default_factory=dict)  # campo -> [(fila, valor, origen)]
    rango_c: str | None = None


class Hoja:
    def __init__(self, ws):
        self.ws = ws
        self.mapa = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.mapa[(r, c)] = rng

    def celda(self, r: int, c: int) -> Celda:
        rng = self.mapa.get((r, c))
        if rng is not None:  # R1
            principal = self.ws.cell(rng.min_row, rng.min_col)
            return Celda(texto(principal.value), principal.coordinate, rng.coord)
        cell = self.ws.cell(r, c)
        return Celda(texto(cell.value), cell.coordinate, None)

    def propia(self, r: int, c: int) -> str | None:
        """Valor escrito en la celda misma (None si es parte no principal de una combinación)."""
        rng = self.mapa.get((r, c))
        if rng is not None and (rng.min_row, rng.min_col) != (r, c):
            return None
        return texto(self.ws.cell(r, c).value)


class Importador:
    def __init__(self, db: Session, ruta: Path, hoja: str):
        self.db, self.ruta, self.nombre_hoja = db, Path(ruta), hoja
        self.obs: list[dict] = []
        self.stats = {"filas_leidas": 0, "filas_servicio": 0, "filas_continuacion": 0, "filas_vacias": 0,
                      "filas_omitidas": 0, "n1": {"creados": 0, "actualizados": 0, "sin_cambios": 0},
                      "n2": {"creados": 0, "actualizados": 0, "sin_cambios": 0}}

    # ------------------------------------------------------------ utilidades
    def observar(self, tipo: str, mensaje: str, rango: str | None = None, codigo: str | None = None, **detalle):
        self.obs.append({"tipo": tipo, "hoja": self.nombre_hoja, "rango": rango, "codigo": codigo,
                         "mensaje": mensaje, "detalle": detalle or None})

    def _abrir(self):
        if not self.ruta.exists():
            raise FileNotFoundError(f"No se encontró el archivo {self.ruta}. Colóquelo en data/CatalogoServicios.xlsx")
        wb = load_workbook(self.ruta, data_only=True)  # solo lectura de valores; el archivo no se modifica
        nombres = {clave(n): n for n in wb.sheetnames}
        if self.nombre_hoja not in wb.sheetnames:
            if clave(self.nombre_hoja) not in nombres:
                raise ValueError(f"La hoja '{self.nombre_hoja}' no existe. Hojas: {wb.sheetnames}")
            self.observar("HOJA_NOMBRE_DIFERENTE", f"Se usó la hoja '{nombres[clave(self.nombre_hoja)]}'.")
            self.nombre_hoja = nombres[clave(self.nombre_hoja)]
        return Hoja(wb[self.nombre_hoja])

    # ------------------------------------------------------------ lectura
    def leer(self, h: Hoja):
        encabezados = [h.celda(FILA_ENCABEZADO, c).valor for c in range(1, 13)]
        for esperado, real, c in zip(COLUMNAS, encabezados, range(1, 13)):
            if clave(esperado) != clave(real or ""):
                self.observar("ENCABEZADO_DIFERENTE", f"Encabezado {get_column_letter(c)}4 = '{real}', se esperaba '{esperado}'.",
                              f"{get_column_letter(c)}4")

        registros: dict[str, Registro] = {}
        anclas: dict[tuple, str] = {}  # (ancla de C) -> código
        n1_nombres: dict[str, list[tuple[int, str, str]]] = {}
        n1_filas: dict[str, list[int]] = {}
        n1_original: dict[str, str] = {}

        for r in range(FILA_INICIO, FILA_FIN + 1):
            self.stats["filas_leidas"] += 1
            celdas = [h.celda(r, c) for c in range(1, 13)]
            propias = [h.propia(r, c) for c in range(1, 13)]
            c_cod = celdas[2]
            if c_cod.valor is None:
                if any(propias):  # R3
                    self.stats["filas_omitidas"] += 1
                    self.observar("FILA_SIN_CODIGO",
                                  f"Fila {r} tiene datos pero no tiene código N2 ni pertenece a un rango combinado de C; se omite.",
                                  f"A{r}:L{r}", valores={CAMPOS[i]: v for i, v in enumerate(propias) if v})
                else:
                    self.stats["filas_vacias"] += 1
                continue

            ancla = c_cod.combinada or c_cod.origen
            cod_raw = str(h.ws[c_cod.origen].value)  # R6: original exacto, sin recortar
            cod = normalizar_codigo(cod_raw)
            if ancla in anclas:
                self.stats["filas_continuacion"] += 1
            else:
                self.stats["filas_servicio"] += 1
                anclas[ancla] = cod
                if cod in registros:
                    self.observar("CODIGO_N2_DUPLICADO",
                                  f"El código {cod} aparece en otro bloque (fila {r}); se integra al registro existente.",
                                  f"C{r}", cod)
                else:
                    registros[cod] = Registro(cod, cod_raw, rango_c=c_cod.combinada or c_cod.origen)
            reg = registros[cod]
            reg.filas.append(r)
            for i, campo in enumerate(CAMPOS):
                if celdas[i].valor is not None:
                    reg.valores.setdefault(campo, []).append((r, celdas[i].valor, celdas[i].origen))

            a = celdas[0]
            if a.valor is not None:
                c1 = normalizar_codigo(a.valor)
                n1_original.setdefault(c1, str(h.ws[a.origen].value))
                n1_filas.setdefault(c1, []).append(r)
                if celdas[1].valor is not None:
                    n1_nombres.setdefault(c1, []).append((r, celdas[1].valor, celdas[1].origen))

        for r in range(FILA_FIN + 1, OPCIONES[0]):
            propias = [h.propia(r, c) for c in range(1, 13)]
            if any(propias) and not any(propias[:4] + propias[8:]) and all(clave(v or "") == "opciones" for v in propias[4:8] if v):
                self.observar("ENCABEZADO_LISTA_OPCIONES", f"Fila {r}: encabezado de las listas de opciones; no es un servicio.",
                              f"E{r}:H{r}")
            elif any(propias):
                self.observar("FILA_FUERA_DE_RANGO", f"La fila {r} tiene contenido fuera del rango de datos (5–101); no se importa.",
                              f"A{r}:L{r}")
        return registros, n1_nombres, n1_filas, n1_original

    def leer_opciones(self, h: Hoja) -> dict[str, list[str]]:
        f0, f1, c0, c1 = OPCIONES
        res = {}
        for c, campo in zip(range(c0, c1 + 1), ["activo", "clase", "criticidad", "tipo"]):
            vals = []
            for r in range(f0, f1 + 1):
                v = h.celda(r, c).valor
                if v and clave(v) != clave(COLUMNAS[c - 1]) and v not in vals:
                    vals.append(v)
            res[campo] = vals
        return res

    # ------------------------------------------------------------ catálogos
    def _catalogos(self):
        return {"clase": ClaseServicio, "criticidad": Criticidad, "tipo": TipoServicio}

    def sincronizar_opciones(self, opciones: dict[str, list[str]]):
        for campo, modelo in self._catalogos().items():
            existentes = {clave(o.nombre): o for o in self.db.scalars(select(modelo))}
            for v in opciones.get(campo, []):
                if clave(v) in existentes:
                    continue
                obj = modelo(nombre=v, activo=True)
                if modelo is Criticidad:
                    obj.orden = len(existentes) + 1
                self.db.add(obj)
                existentes[clave(v)] = obj
                self.observar("OPCION_AGREGADA", f"La lista de opciones contiene '{v}' en {campo}; se agregó al catálogo.",
                              "E112:H122", valor=v, campo=campo)
        self.db.flush()

    def resolver(self, campo: str, valor: str | None, codigo: str):
        """Devuelve (id, original) según R7. Registra mapeos de etiqueta si se corrigió la forma."""
        if valor is None:
            return None, None
        modelo = self._catalogos()[campo]
        exacto = self.db.scalar(select(modelo).where(modelo.nombre == valor))
        if exacto:
            return exacto.id, valor
        for o in self.db.scalars(select(modelo)):
            if clave(o.nombre) == clave(valor):
                self._mapear(campo, valor, o.nombre, "Diferencia de mayúsculas/espacios/tildes", codigo)
                return o.id, valor
        self.observar("VALOR_FUERA_DE_CATALOGO", f"{campo.upper()} = '{valor}' no está en el catálogo; se conserva como original.",
                      None, codigo, campo=campo, valor=valor)
        return None, valor

    def _mapear(self, campo: str, original: str, canonico: str, motivo: str, codigo: str | None = None):
        m = self.db.scalar(select(MapeoEtiqueta).where(MapeoEtiqueta.campo == campo, MapeoEtiqueta.valor_original == original))
        if m is None:
            self.db.add(MapeoEtiqueta(campo=campo, valor_original=original, valor_canonico=canonico, motivo=motivo))
            self.db.flush()
        self.observar("ETIQUETA_MAPEADA", f"{campo.upper()} '{original}' se asoció al valor de catálogo '{canonico}' ({motivo}).",
                      None, codigo, campo=campo, original=original, canonico=canonico)

    # ------------------------------------------------------------ escritura
    @staticmethod
    def _hash(d: dict) -> str:
        return hashlib.sha256(json.dumps(d, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()

    def _primero(self, reg: Registro, campo: str) -> str | None:
        vals = reg.valores.get(campo, [])
        distintos = []
        for fila, v, origen in vals:
            if v not in [d[1] for d in distintos]:
                distintos.append((fila, v, origen))
        if len(distintos) > 1 and campo not in ("cod_n1", "nombre_n1", "cod_n2"):  # R5
            self.observar("CONFLICTO_ATRIBUTO",
                          f"{reg.codigo}: el campo '{campo}' tiene valores distintos en filas del mismo servicio; "
                          f"se conserva '{distintos[0][1]}' (fila {distintos[0][0]}).",
                          reg.rango_c, reg.codigo, valores=[{"fila": f, "valor": v, "celda": o} for f, v, o in distintos])
        return distintos[0][1] if distintos else None

    def guardar_n1(self, n1_nombres, n1_filas, n1_original) -> dict[str, ServicioN1]:
        resultado = {}
        for cod, filas in sorted(n1_filas.items()):
            nombres = []
            for fila, v, origen in n1_nombres.get(cod, []):
                if v not in [n["valor"] for n in nombres]:
                    nombres.append({"fila": fila, "celda": origen, "valor": v})
            rango = f"A{min(filas)}:B{max(filas)}"
            if not nombres:
                self.observar("N1_SIN_NOMBRE", f"El servicio N1 {cod} no tiene nombre; se usa el código.", rango, cod)
                nombres = [{"fila": min(filas), "celda": None, "valor": cod}]
            canonico = NOMBRE_CANONICO_N1.get(cod) if cod in NOMBRE_CANONICO_N1 else nombres[0]["valor"]
            if canonico not in [n["valor"] for n in nombres]:
                canonico = nombres[0]["valor"]
            if len(nombres) > 1:  # R4
                self.observar("CONFLICTO_NOMBRE_N1",
                              f"El código N1 {cod} aparece con {len(nombres)} nombres distintos; nombre canónico: '{canonico}'.",
                              rango, cod, nombres=nombres, regla="NOMBRE_CANONICO_N1" if cod in NOMBRE_CANONICO_N1 else "primera aparición")
            original = n1_original[cod]
            if original.strip() != cod:
                self.observar("CODIGO_NORMALIZADO", f"Código N1 '{original}' normalizado a '{cod}'.", rango, cod, original=original)
            datos = {"codigo": cod, "codigo_original": original, "nombre": canonico,
                     "nombres_alternativos": nombres, "origen_rango": rango}
            h = self._hash(datos)
            obj = self.db.scalar(select(ServicioN1).where(ServicioN1.codigo == cod))
            if obj is None:
                obj = ServicioN1(activo=True, origen_hoja=self.nombre_hoja, origen_hash=h, **datos)
                self.db.add(obj)
                self.stats["n1"]["creados"] += 1
            elif obj.origen_hash != h:
                for k, v in datos.items():
                    setattr(obj, k, v)
                obj.origen_hoja, obj.origen_hash = self.nombre_hoja, h
                self.stats["n1"]["actualizados"] += 1
            else:
                self.stats["n1"]["sin_cambios"] += 1
            resultado[cod] = obj
        self.db.flush()
        return resultado

    def guardar_n2(self, registros: dict[str, Registro], n1s: dict[str, ServicioN1]):
        observados = set()
        for cod, reg in registros.items():
            antes = len(self.obs)
            rango = f"A{min(reg.filas)}:L{max(reg.filas)}"
            cod_n1_raw = self._primero(reg, "cod_n1")
            distintos_n1 = {normalizar_codigo(v) for _, v, _ in reg.valores.get("cod_n1", [])}
            if len(distintos_n1) > 1:
                self.observar("CONFLICTO_N1_EN_SERVICIO", f"{cod} abarca filas con distintos códigos N1 {sorted(distintos_n1)}; se usa el primero.",
                              rango, cod)
            if cod_n1_raw is None:
                # R9: sin COD.N1 en la fila ni en un rango combinado. Se deriva del prefijo jerárquico
                # del propio código N2 (SE.12.3 -> SE.12) solo si ese N1 existe; no se copia de filas vecinas.
                prefijo = cod.rsplit(".", 1)[0] if "." in cod else None
                if prefijo not in n1s:
                    self.observar("N2_SIN_N1", f"{cod} no tiene código de nivel 1 y no se puede derivar; se omite.", rango, cod)
                    self.stats["filas_omitidas"] += len(reg.filas)
                    continue
                self.observar("N1_DERIVADO_DEL_CODIGO",
                              f"{cod} no tiene COD.N1 en el Excel; se asocia a {prefijo} por el prefijo de su código.",
                              rango, cod, regla="R9 prefijo del código N2", n1=prefijo)
                cod_n1_raw = prefijo
            n1 = n1s[normalizar_codigo(cod_n1_raw)]
            if reg.codigo_original.strip() != cod:
                self.observar("CODIGO_NORMALIZADO", f"Código N2 '{reg.codigo_original}' normalizado a '{cod}'.", rango, cod,
                              original=reg.codigo_original)

            nombre = self._primero(reg, "nombre_n2")
            if nombre is None:
                self.observar("ATRIBUTO_AUSENTE", f"{cod} no tiene nombre de nivel 2; se usa el código.", rango, cod)
                nombre = cod
            activo_orig = self._primero(reg, "activo")
            activo = ACTIVO_ALIAS.get(clave(activo_orig).upper()) if activo_orig else None
            if activo_orig and activo is None:
                self.observar("VALOR_FUERA_DE_CATALOGO", f"ACTIVO = '{activo_orig}' no es S/N; se conserva como desconocido.",
                              rango, cod, campo="activo", valor=activo_orig)
            elif activo_orig and activo != activo_orig:
                self._mapear("activo", activo_orig, activo, "Equivalencia explícita de ACTIVO", cod)
            clase_id, clase_o = self.resolver("clase", self._primero(reg, "clase"), cod)
            crit_id, crit_o = self.resolver("criticidad", self._primero(reg, "criticidad"), cod)
            tipo_id, tipo_o = self.resolver("tipo", self._primero(reg, "tipo"), cod)
            metrica = self._primero(reg, "metrica")
            minimo, min_o = self._numero(self._primero(reg, "minimo"), cod, "minimo", rango)
            maximo, max_o = self._numero(self._primero(reg, "maximo"), cod, "maximo", rango)
            if minimo is not None and maximo is not None and minimo > maximo:
                self.observar("UMBRAL_INVERTIDO", f"{cod}: mínimo {minimo} > máximo {maximo}; se conservan solo como originales.",
                              rango, cod)
                minimo = maximo = None

            ausentes = [n for n, v in (("activo", activo_orig), ("clase", clase_o), ("criticidad", crit_o),
                                       ("tipo", tipo_o), ("metrica", metrica)) if v is None]
            if ausentes:  # R7
                self.observar("ATRIBUTO_AUSENTE", f"{cod} no tiene valor en: {', '.join(ausentes)}. Queda como desconocido.",
                              rango, cod, campos=ausentes)
            informativas = {"ETIQUETA_MAPEADA", "CODIGO_NORMALIZADO"}
            pendientes = [o for o in self.obs[antes:] if o["tipo"] not in informativas]
            revision = "REVISION" if (None in (activo, clase_id, crit_id, tipo_id)) or pendientes else "OK"

            origen = {campo: [{"fila": f, "valor": v, "celda": o} for f, v, o in vals] for campo, vals in reg.valores.items()}
            datos = {
                "servicio_n1_id": n1.id, "codigo": cod, "codigo_original": reg.codigo_original, "nombre": nombre,
                "activo": activo, "activo_original": activo_orig,
                "clase_id": clase_id, "clase_original": clase_o, "criticidad_id": crit_id, "criticidad_original": crit_o,
                "tipo_id": tipo_id, "tipo_original": tipo_o, "descripcion": self._primero(reg, "descripcion"),
                "metrica": metrica, "minimo": minimo, "minimo_original": min_o, "maximo": maximo, "maximo_original": max_o,
                "estado_revision": revision, "origen_rango": rango, "origen_valores": origen,
            }
            h = self._hash(datos)
            obj = self.db.scalar(select(ServicioN2).where(ServicioN2.codigo == cod))
            if obj is None:
                self.db.add(ServicioN2(origen_hoja=self.nombre_hoja, origen_hash=h, dado_de_baja=False, **datos))
                self.stats["n2"]["creados"] += 1
            elif obj.origen_hash != h:
                for k, v in datos.items():  # seccion, responsable y baja lógica NO se tocan
                    setattr(obj, k, v)
                obj.origen_hoja, obj.origen_hash = self.nombre_hoja, h
                self.stats["n2"]["actualizados"] += 1
            else:
                self.stats["n2"]["sin_cambios"] += 1
            if len(self.obs) > antes:
                observados.add(cod)
        self.db.flush()
        return observados

    def _numero(self, v: str | None, cod: str, campo: str, rango: str):
        if v is None:
            return None, None  # ausente se mantiene ausente, nunca 0
        try:
            return Decimal(v.replace(",", ".")), v
        except InvalidOperation:
            self.observar("VALOR_NO_NUMERICO", f"{cod}: {campo} = '{v}' no es numérico; se conserva como original.", rango, cod)
            return None, v

    # ------------------------------------------------------------ ejecución
    def ejecutar(self) -> dict:
        sha_antes = hashlib.sha256(self.ruta.read_bytes()).hexdigest() if self.ruta.exists() else None
        h = self._abrir()
        registros, n1_nombres, n1_filas, n1_original = self.leer(h)
        self.sincronizar_opciones(self.leer_opciones(h))
        n1s = self.guardar_n1(n1_nombres, n1_filas, n1_original)
        observados = self.guardar_n2(registros, n1s)
        sha_despues = hashlib.sha256(self.ruta.read_bytes()).hexdigest()
        if sha_antes != sha_despues:
            raise RuntimeError("El archivo Excel cambió durante la importación")

        n1_codigos = {normalizar_codigo(r.valores["cod_n1"][0][1]) for r in registros.values() if r.valores.get("cod_n1")}
        control = {"n1_encontrados": len(n1_codigos), "n2_encontrados": len(registros),
                   "esperado_n1": config.ESPERADO_N1, "esperado_n2": config.ESPERADO_N2}
        control["ok"] = control["n1_encontrados"] == config.ESPERADO_N1 and control["n2_encontrados"] == config.ESPERADO_N2
        if not control["ok"]:
            self.observar("CONTROL_NO_COINCIDE",
                          f"Se encontraron {control['n1_encontrados']} N1 y {control['n2_encontrados']} N2; "
                          f"se esperaban {config.ESPERADO_N1} y {config.ESPERADO_N2}.")
        resumen = {**self.stats, "control": control, "observaciones": len(self.obs),
                   "servicios_observados": sorted(observados),
                   "observaciones_por_tipo": {}}
        for o in self.obs:
            resumen["observaciones_por_tipo"][o["tipo"]] = resumen["observaciones_por_tipo"].get(o["tipo"], 0) + 1
        imp = Importacion(archivo=self.ruta.name, sha256=sha_despues, hoja=self.nombre_hoja, resumen=resumen)
        self.db.add(imp)
        self.db.flush()
        for o in self.obs:
            self.db.add(ObservacionImportacion(importacion_id=imp.id, **o))
        self.db.commit()
        resumen["importacion_id"] = imp.id
        return resumen


def importar(db: Session, ruta: Path | str | None = None, hoja: str | None = None) -> dict:
    try:
        return Importador(db, Path(ruta or config.EXCEL_PATH), hoja or config.EXCEL_HOJA).ejecutar()
    except Exception:
        db.rollback()
        raise
